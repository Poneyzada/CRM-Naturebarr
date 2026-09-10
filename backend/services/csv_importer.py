import csv
import io
import re
from typing import List, Dict, Any, Tuple
import openpyxl

def normalize_key(k: str) -> str:
    """Normaliza nome da coluna removendo acentos, pontuação e espaços."""
    s = k.strip().lower()
    s = re.sub(r'[áàãâä]', 'a', s)
    s = re.sub(r'[éèêë]', 'e', s)
    s = re.sub(r'[íìîï]', 'i', s)
    s = re.sub(r'[óòõôö]', 'o', s)
    s = re.sub(r'[úùûü]', 'u', s)
    s = re.sub(r'[ç]', 'c', s)
    s = re.sub(r'[^a-z0-9]', '_', s)
    s = re.sub(r'_+', '_', s).strip('_')
    return s

def detect_region(cidade: str, estado: str) -> str:
    """Detecta automaticamente a região comercial (Salvador, SP, SC ou Outro)."""
    cid = cidade.lower() if cidade else ""
    uf = estado.upper() if estado else ""

    if "salvador" in cid or "lauro de freitas" in cid or "camacari" in cid or uf == "BA":
        return "Salvador"
    if "sao paulo" in cid or "são paulo" in cid or "campinas" in cid or "santos" in cid or "ribeirao" in cid or uf == "SP":
        return "Sao Paulo"
    if "florianopolis" in cid or "florianópolis" in cid or "joinville" in cid or "blumenau" in cid or "balneario" in cid or "itajai" in cid or uf == "SC":
        return "Santa Catarina"
    return "Outro"

def parse_csv_or_excel(file_bytes: bytes, filename: str) -> Tuple[List[Dict[str, Any]], List[str]]:
    """Lê bytes de CSV ou XLSX e retorna lista de dicionários mapeados e erros."""
    rows = []
    errors = []

    try:
        if filename.endswith(".xlsx") or filename.endswith(".xls"):
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            sheet = wb.active
            iter_rows = sheet.iter_rows(values_only=True)
            headers = next(iter_rows, None)
            if not headers:
                return [], ["Arquivo Excel vazio"]
            
            raw_headers = [str(h or "") for h in headers]
            for row_idx, row in enumerate(iter_rows, start=2):
                row_dict = {}
                for h, val in zip(raw_headers, row):
                    if h:
                        row_dict[h] = str(val).strip() if val is not None else ""
                if any(row_dict.values()):
                    rows.append(row_dict)
        else:
            # Tenta decodificar como utf-8, se falhar tenta latin-1
            try:
                content = file_bytes.decode("utf-8-sig")
            except UnicodeDecodeError:
                content = file_bytes.decode("latin-1")

            # Detecta delimitador (, ou ;)
            first_line = content.splitlines()[0] if content.splitlines() else ""
            delimiter = ";" if first_line.count(";") > first_line.count(",") else ","

            reader = csv.DictReader(io.StringIO(content), delimiter=delimiter)
            for r in reader:
                rows.append({k: (v.strip() if v else "") for k, v in r.items() if k})

    except Exception as e:
        return [], [f"Erro ao processar arquivo: {str(e)}"]

    # Mapeamento inteligente de colunas
    mapped_pdvs = []
    for idx, r in enumerate(rows, start=1):
        norm_row = {normalize_key(k): v for k, v in r.items()}

        # Identifica Razão Social e Nome Fantasia
        razao = (
            norm_row.get("razao_social") or
            norm_row.get("razao") or
            norm_row.get("empresa") or
            norm_row.get("nome_da_empresa") or
            norm_row.get("nome_fantasia") or
            norm_row.get("fantasia") or
            norm_row.get("pdv") or
            norm_row.get("loja") or
            f"PDV Importado #{idx}"
        )
        fantasia = (
            norm_row.get("nome_fantasia") or
            norm_row.get("fantasia") or
            norm_row.get("loja") or
            norm_row.get("pdv") or
            razao
        )

        # Telefone / WhatsApp
        tel = (
            norm_row.get("telefone_whatsapp") or
            norm_row.get("whatsapp") or
            norm_row.get("celular") or
            norm_row.get("telefone") or
            norm_row.get("tel") or
            norm_row.get("contato_whatsapp") or
            ""
        )
        if not tel:
            # Se não tiver telefone, coloca um placeholder para não falhar a importação de cadastro
            tel = "71999999999"

        # CNPJ/CPF
        cnpj = (
            norm_row.get("cnpj_cpf") or
            norm_row.get("cnpj") or
            norm_row.get("cpf") or
            norm_row.get("documento") or
            None
        )

        # Responsável
        resp = (
            norm_row.get("responsavel") or
            norm_row.get("comprador") or
            norm_row.get("contato") or
            norm_row.get("nome") or
            norm_row.get("dono") or
            "Responsável"
        )

        # Localização
        cid = (
            norm_row.get("cidade") or
            norm_row.get("municipio") or
            "Salvador"
        )
        uf = (
            norm_row.get("estado") or
            norm_row.get("uf") or
            ("BA" if "salvador" in cid.lower() else "SP")
        )
        regiao = (
            norm_row.get("regiao") or
            norm_row.get("polo") or
            detect_region(cid, uf)
        )

        # Categoria
        cat = norm_row.get("categoria") or norm_row.get("segmento") or norm_row.get("tipo") or "Emporio"
        # Ajusta categoria para valores válidos
        valid_cats = ["Emporio", "Academia", "Box Crossfit", "Farmacia", "Suplementos", "Cafeteria", "Mercado Saudavel", "Outro"]
        matched_cat = "Emporio"
        for vc in valid_cats:
            if vc.lower() in cat.lower():
                matched_cat = vc
                break

        mapped_pdvs.append({
            "razao_social": razao,
            "nome_fantasia": fantasia,
            "cnpj_cpf": cnpj,
            "responsavel": resp,
            "telefone_whatsapp": tel,
            "email": norm_row.get("email") or norm_row.get("e_mail"),
            "endereco": norm_row.get("endereco") or norm_row.get("logradouro"),
            "bairro": norm_row.get("bairro"),
            "cidade": cid,
            "estado": uf[:2].upper() if uf else "BA",
            "regiao": regiao if regiao in ["Salvador", "Sao Paulo", "Santa Catarina", "Outro"] else "Salvador",
            "cep": norm_row.get("cep"),
            "categoria": matched_cat,
            "origem": "Importacao Planilha",
            "status_pipeline": "PDV Ativo" if norm_row.get("comprando") in ["sim", "true", "1", "ativo"] else "Novo Lead",
            "media_dias_recompra": 21
        })

    return mapped_pdvs, errors
